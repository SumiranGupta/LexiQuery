"""
Document Loaders
Unified interface for loading documents from multiple formats.
Extensible architecture — add new loaders by registering in LOADER_MAP.
"""

from pathlib import Path
from typing import List

from langchain_core.documents import Document
from langchain_community.document_loaders import (
    PDFPlumberLoader,
    TextLoader,
    CSVLoader,
    Docx2txtLoader,
)

from backend.utils.logger import get_logger
from backend.utils.config import SUPPORTED_EXTENSIONS

logger = get_logger("ingestion.loaders")

# Registry of file extension → loader class (None = custom loader used below)
LOADER_MAP = {
    ".pdf": PDFPlumberLoader,
    ".txt": TextLoader,
    ".csv": CSVLoader,
    ".md": TextLoader,
    ".docx": Docx2txtLoader,
    ".pptx": None,   # handled by _load_pptx
    ".xlsx": None,   # handled by _load_xlsx
}


def _load_pptx(file_path: Path) -> List[Document]:
    """Load a PPTX file slide-by-slide using python-pptx for full text extraction."""
    from pptx import Presentation as PptxPresentation

    prs = PptxPresentation(str(file_path))
    documents = []

    for slide_num, slide in enumerate(prs.slides, 1):
        parts: List[str] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        parts.append(" | ".join(row_text))

        content = "\n".join(parts)
        if content.strip():
            documents.append(Document(
                page_content=content,
                metadata={
                    "source": str(file_path),
                    "page": slide_num,
                    "total_slides": len(prs.slides),
                },
            ))

    return documents


def _load_xlsx(file_path: Path) -> List[Document]:
    """Load an Excel file sheet-by-sheet using openpyxl for reliable text extraction."""
    import openpyxl

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    documents = []

    for sheet_num, sheet_name in enumerate(wb.sheetnames, 1):
        ws = wb[sheet_name]
        rows_text: List[str] = []

        # Extract header row separately for context
        header = []
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            header = [str(c).strip() for c in row if c is not None]

        for row in ws.iter_rows(min_row=2, values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                if header:
                    # Pair header with value for richer semantic context
                    paired = "; ".join(
                        f"{h}: {v}" for h, v in zip(header, [str(c).strip() for c in row if c is not None])
                        if v.strip()
                    )
                    if paired:
                        rows_text.append(paired)
                else:
                    rows_text.append(" | ".join(cells))

        content = f"Sheet: {sheet_name}\n" + "\n".join(rows_text)
        if rows_text:
            documents.append(Document(
                page_content=content,
                metadata={
                    "source": str(file_path),
                    "page": sheet_num,
                    "sheet_name": sheet_name,
                    "total_sheets": len(wb.sheetnames),
                },
            ))

    return documents


def load_document(file_path: str | Path) -> List[Document]:
    """Load a document file and return LangChain Documents with enriched metadata."""
    file_path = Path(file_path)
    ext = file_path.suffix.lower()

    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {SUPPORTED_EXTENSIONS}")

    if ext not in LOADER_MAP:
        raise ValueError(f"No loader registered for extension: {ext}")

    logger.info(f"Loading document: {file_path.name} ({ext})")

    try:
        # Use custom loader for PPTX and XLSX
        if ext == ".pptx":
            documents = _load_pptx(file_path)
        elif ext == ".xlsx":
            documents = _load_xlsx(file_path)
        else:
            loader_cls = LOADER_MAP[ext]
            loader = loader_cls(str(file_path))
            documents = loader.load()

        # Enrich metadata for every page/section
        for i, doc in enumerate(documents):
            doc.metadata.update({
                "source_file": file_path.name,
                "file_type": ext,
                "page": doc.metadata.get("page", i + 1),
            })

        logger.info(f"Loaded {len(documents)} sections from {file_path.name}")
        return documents

    except Exception as e:
        logger.error(f"Failed to load {file_path.name}: {e}")
        raise
